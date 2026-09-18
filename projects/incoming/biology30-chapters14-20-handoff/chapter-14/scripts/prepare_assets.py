from pathlib import Path
import fitz,json,hashlib,shutil,html,os
ROOT=Path(__file__).resolve().parents[1]; W=ROOT/'workspace'; A=W/'assets';

# Rebuilding uses the packaged chapter PDF, not an external Downloads path.
P=A/'textbook/chapter-14.pdf';P.parent.mkdir(exist_ok=True,parents=True)
if not P.exists():raise FileNotFoundError('Restore workspace/assets/textbook/chapter-14.pdf before rebuilding.')
pdf=fitz.open(P);hashpdf=hashlib.sha256(P.read_bytes()).hexdigest()
for d in ['figures','labeling','textbook-practice','assessment']: (A/d).mkdir(exist_ok=True)
records=[];labels=[]
def source_crop(id,page,box,title,caption,alt,lesson):
 path=A/'figures'/f'{id}.png';p=pdf[page-476];p.get_pixmap(matrix=fitz.Matrix(3,3),clip=fitz.Rect(box)).save(path)
 rec={'id':id,'title':title,'lesson':lesson,'src':'./assets/figures/'+path.name,'caption':caption,'alt':alt,'kind':'source-crop','source':{'pdf':'assets/textbook/chapter-14.pdf','sha256':hashpdf,'printedPage':page,'pdfPage':page-475,'box':box},'reviewStatus':'Source location and intended structures inspected; final rendered crop inspection recorded separately.'};records.append(rec);return rec
source_crop('male-anatomy',479,[65,490,442,742],'Male reproductive structures','Follow the connected ducts from the testis to the urethra. Nearby organs are not automatically part of that route. Inquiry into Biology, p. 479.','Side view of male reproductive anatomy with labelled testes, epididymis, vas deferens, accessory glands and urethra.',2)
source_crop('testis-cells',480,[160,48,543,378],'Inside the testis','Developing sperm and Sertoli cells occupy the tubules; interstitial cells lie between tubules. Inquiry into Biology, p. 480.','Testis, seminiferous tubules, interstitial tissue and enlarged developing sperm with supporting Sertoli cells.',3)
source_crop('sperm-structure',480,[165,240,377,384],'The structure of a sperm','The acrosome, head, middle piece and flagellum have different roles. Inquiry into Biology, p. 480.','A sperm with labelled acrosome, head, mitochondrial middle section and long flagellum.',3)
source_crop('female-anatomy',482,[158,405,539,742],'Female reproductive structures','Locate the oviducts, uterus, cervix and vagina before tracing either gamete’s route. Inquiry into Biology, p. 482.','Side view of female reproductive anatomy showing the ovary and oviduct near the uterus, cervix and vagina.',5)
source_crop('follicle-development',496,[160,416,535,740],'A follicle before and after ovulation','An oocyte leaves the mature follicle; hormone-producing tissue remains in the ovary. Inquiry into Biology, p. 496.','Developing follicles, a mature follicle, released oocyte and corpus luteum in an enlarged ovary.',6)
# Change only printed structure labels. Existing source leader lines stay in place.
def relabel(id,srcid,lesson,box,items, title,alt):
 rec=next(x for x in records if x['id']==srcid);page=rec['source']['printedPage'];doc=fitz.open(P);p=doc[page-476]
 for letter,answer,rect in items:
  # A white rectangle covers only the source text; no anatomical pixels/endpoints are moved.
  r=fitz.Rect(rect);p.draw_rect(r,color=None,fill=(1,1,1),overlay=True)
  p.insert_text((r.x0+1,r.y1-1),letter,fontsize=11,fontname='hebo',color=(0,0,0),overlay=True)
 out=A/'labeling'/f'{id}.png';p.get_pixmap(matrix=fitz.Matrix(3,3),clip=fitz.Rect(box)).save(out);doc.close()
 d={'id':id,'title':title,'lesson':lesson,'src':'./assets/labeling/'+out.name,'alt':alt,'answers':{a:b for a,b,_ in items},'options':list(dict.fromkeys(b for _,b,_ in items)),'reviewNote':'Letter positions replace source text beside unchanged source leader lines.'}
 labels.append(d);records.append({'id':id,'kind':'source-label-adaptation','sourceFigure':srcid,'lesson':lesson,'src':d['src'],'source':rec['source'],'crop':box,'changes':[{'letter':a,'answer':b,'textCoverBox':r}for a,b,r in items],'reviewStatus':'Requires final crop inspection before reported visual pass.'})
relabel('ch14-label-male','male-anatomy',2,[65,490,442,742],[
 ('A','Vas deferens',[71,612,134,624]),('B','Urethra',[71,634,102,646]),('C','Penis',[71,656,95,668]),
 ('D','Seminal vesicle',[377,632,436,644]),('E','Ejaculatory duct',[377,647,440,659]),('F','Prostate gland',[377,662,436,674]),
 ('G','Bulbourethral gland',[377,676,438,688]),('H','Epididymis',[377,691,422,703]),('I','Testis',[377,705,400,717]),('J','Scrotum',[377,720,412,732])],
 'Male reproductive anatomy','Lettered male reproductive anatomy; follow the original leader lines. Urinary and other orientation structures retain their names.')
relabel('ch14-label-sperm','sperm-structure',3,[165,240,377,384],[
 ('A','Acrosome',[216,251,257,263]),('B','Head',[229,264,253,276]),('C','Middle piece',[174,298,207,321]),('D','Flagellum',[198,336,257,349])],
 'Sperm structure','A source sperm drawing with four original label positions replaced by A through D.')
# Follicle-only crop removes surrounding explanatory captions that would disclose answers.
relabel('ch14-label-follicle','follicle-development',6,[160,490,535,714],[
 ('A','Developing follicles',[354,498,431,511]),('B','Mature follicle',[460,523,518,536]),('C','Oocyte',[494,573,523,586]),
 ('D','Released oocyte',[496,619,532,643]),('E','Corpus luteum',[366,706,425,718])],
 'Follicle development','An enlarged ovary with lettered developing follicles, mature follicle, oocyte, released cell and corpus-luteum region.')
# Original course schematics use explicit relationships rather than approximate anatomy.
SVG_HEAD='''<svg xmlns="http://www.w3.org/2000/svg" width="1200" height="820" viewBox="0 0 1200 820"><rect width="1200" height="820" fill="white"/><defs><marker id="a" markerWidth="10" markerHeight="10" refX="8" refY="5" orient="auto"><path d="M0 0 L10 5 L0 10 Z" fill="#146c60"/></marker></defs><style>text{font-family:Arial,sans-serif;fill:#171b1b}.title{font-size:32px;font-weight:700}.body{font-size:22px}.small{font-size:18px}.letter{font-size:30px;font-weight:700}</style>'''
def text(x,y,s,cls='body',anchor='middle'):return f'<text x="{x}" y="{y}" class="{cls}" text-anchor="{anchor}">{html.escape(s)}</text>'
def box(x,y,w,h,lines,fill='#eef3ee'):
 out=f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="8" fill="{fill}" stroke="#146c60" stroke-width="2"/>'
 for i,line in enumerate(lines):out+=text(x+w/2,y+h/2+(i-(len(lines)-1)/2)*25+7,line,'letter' if len(line)==1 else 'body')
 return out
def arrow(x1,y1,x2,y2):return f'<path d="M{x1} {y1} L{x2} {y2}" stroke="#146c60" stroke-width="4" fill="none" marker-end="url(#a)"/>'
def save_svg(id,title,lesson,svg,caption,alt,labelmap=None):
 folder='labeling' if labelmap else 'figures';path=A/folder/f'{id}.svg';path.write_text(svg+'</svg>')
 if labelmap: labels.append({'id':id,'title':title,'lesson':lesson,'src':f'./assets/{folder}/{id}.svg','alt':alt,'answers':labelmap,'options':list(labelmap.values()),'reviewNote':'Authored relationship model, not an anatomical scale drawing.'})
 records.append({'id':id,'title':title,'lesson':lesson,'src':f'./assets/{folder}/{id}.svg','caption':caption,'alt':alt,'kind':'authored-schematic','source':{'basis':'Chapter 14 pp. 493–498 and teacher notes; reconciled in source-discrepancies.json'},'reviewStatus':'Relationships and answer identities checked programmatically; full-size inspection recorded separately.'})
for blank in [False,True]:
 s=SVG_HEAD+text(600,48,'Hormonal control of sperm production','title')
 s+=box(405,80,390,64,['A']if blank else['Hypothalamus'])+arrow(600,145,600,185)
 s+=box(495,194,210,55,['B']if blank else['GnRH'])+arrow(600,251,600,290)
 s+=box(405,301,390,66,['C']if blank else['Anterior pituitary'])
 s+=arrow(500,368,300,420)+arrow(700,368,900,420)
 s+=box(190,432,220,58,['D']if blank else['FSH'])+box(790,432,220,58,['E']if blank else['LH'])
 s+=arrow(300,492,300,535)+arrow(900,492,900,535)
 s+=box(135,548,330,72,['F']if blank else['Sertoli cells','in seminiferous tubules'])+box(735,548,330,72,['G']if blank else['Interstitial cells','between tubules'])
 s+=arrow(900,621,900,662)+box(790,671,220,55,['H']if blank else['Testosterone'])
 s+=arrow(300,621,300,677)+box(135,690,330,60,['I']if blank else['Support sperm production'])
 s+='<path d="M790 698 L465 718" stroke="#146c60" stroke-width="4" marker-end="url(#a)"/>'
 s+='<path d="M1010 698 L1115 698 L1115 110 L804 110 M1115 332 L804 332" stroke="#a43f35" stroke-width="3" fill="none" stroke-dasharray="8 6"/><path d="M804 95 L804 125 M804 317 L804 347" stroke="#a43f35" stroke-width="4"/>'
 s+=text(1129,490,'−','letter')+text(600,790,'Dashed return lines: testosterone inhibits upstream stimulation.','small')
 save_svg('ch14-label-feedback'if blank else'male-feedback','Male hormonal feedback',4,s,'Follow release sites, hormones and target cells separately. Dashed lines show negative feedback.','Two pituitary branches act on different testicular cells; testosterone feeds back to the control centres.',{'A':'Hypothalamus','B':'GnRH','C':'Anterior pituitary','D':'FSH','E':'LH','F':'Sertoli cells','G':'Interstitial cells','H':'Testosterone','I':'Support of sperm production'}if blank else None)
# A paired event timeline, explicitly schematic, using exact course palette.
s=SVG_HEAD+text(600,48,'Ovarian and uterine events occur together','title')
s+=text(600,82,'Illustrative 28-day model — actual cycle timing varies','small')
for x,top,bot in [(70,['Early follicular phase','Follicles develop'],['Menstrual flow','Part of lining is shed']),(430,['Late follicular phase','Estrogen rises'],['Rebuilding','Endometrium grows']),(790,['After ovulation','Corpus luteum active'],['Maintenance','Progesterone supports lining'])]:
 s+=box(x,210,320,125,top)+box(x,485,320,125,bot)+arrow(x+160,340,x+160,478)
s+=text(55,175,'OVARY','small','start')+text(55,453,'UTERUS','small','start')+arrow(395,271,423,271)+arrow(755,271,783,271)
s+=text(600,680,'Flow overlaps early follicular development.','body')+text(600,724,'Hormone decline after corpus-luteum regression permits the next flow.','small')
save_svg('cycle-coordination','Coordinating two cycles',7,s,'Flow is a uterine event that overlaps the early ovarian follicular phase. This is a sequence model, not a personal calendar.','Paired ovarian and uterine stages showing simultaneous events and hormonal links.')
# Qualitative graph: numeric samples are rendering controls, NOT observed patient data.
import math
xs=list(range(1,29)); vals={
 'Estrogen':[.12+.72*math.exp(-((d-12)/3.2)**2)+.3*math.exp(-((d-21)/4.5)**2)for d in xs],
 'LH':[.10+.86*math.exp(-((d-14)/.8)**2)for d in xs],
 'FSH':[.10+.11*math.exp(-((d-3)/3.2)**2)+.23*math.exp(-((d-14)/1.4)**2)for d in xs],
 'Progesterone':[.04+.70*math.exp(-((d-21)/4.5)**2)if d>=14 else.04 for d in xs]}
colors={'Estrogen':'#146c60','LH':'#a43f35','FSH':'#a15c00','Progesterone':'#154212'}
for blank in [False,True]:
 s=SVG_HEAD+text(600,45,'Hormone patterns across an illustrative cycle','title')+text(600,83,'Schematic relative levels — not measured data; compare timing, not absolute amounts','small')
 for row,names in enumerate([['LH','FSH'],['Estrogen','Progesterone']]):
  top=155+row*285;s+=f'<path d="M110 {top} V{top+220} H1110" stroke="#171b1b" fill="none" stroke-width="2"/>'
  for d in [1,7,14,21,28]:
   x=110+(d-1)/27*970;s+=f'<path d="M{x} {top+220} v8" stroke="#171b1b"/>'+text(x,top+252,str(d),'small')
  for k,name in enumerate(names):
   pts=' '.join(f'{110+(d-1)/27*970:.1f},{top+218-vals[name][d-1]*205:.1f}'for d in xs)
   s+=f'<polyline points="{pts}" fill="none" stroke="{colors[name]}" stroke-width="5"'+(' stroke-dasharray="12 6"'if k else'')+'/>'
   letter=['A','B','C','D'][row*2+k];s+=text(165+k*240,top-20,letter if blank else name,'body')
   s+=f'<path d="M{120+k*240} {top-26} h28" stroke="{colors[name]}" stroke-width="5"'+(' stroke-dasharray="12 6"'if k else'')+'/>'
  s+=text(65,top+95,'Relative','small')+text(65,top+120,'level','small')
 s+=text(600,735,'Cycle day (day 1 = start of menstrual flow)','body')
 s+=text(600,780,'Typical order: estrogen rise → LH surge / ovulation → progesterone rise','small') if not blank else text(600,780,'Identify A–D from the timing and shape of the curves.','small')
 save_svg('ch14-label-cycle'if blank else'cycle-hormones','Identify the hormone patterns',8,s,'An authored qualitative model based on the teaching sequence. Curves are separately scaled and are not patient observations.','Two relative-level plots with a brief pituitary surge and later ovarian hormone change.',{'A':'LH','B':'FSH','C':'Estrogen','D':'Progesterone'}if blank else None)
# Data source remains with authoring, not presented as biological measurements.
(ROOT/'authoring/schematic-rendering-values.json').write_text(json.dumps({'status':'Illustrative drawing controls; not experimental data','cycle':vals},indent=2))
(ROOT/'authoring/diagram-review.json').write_text(json.dumps(records,indent=2));(ROOT/'authoring/label-diagrams.json').write_text(json.dumps(labels,indent=2))
print('Prepared',len(records),'figures;',len(labels),'labeling activities')
# Final isolation pass: remove neighbouring figure/text fragments outside the intended drawing.
from PIL import Image,ImageDraw,ImageFont
font=ImageFont.truetype(os.environ.get('BIOLOGY_LABEL_FONT','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'),22)
def clean_png(path,sourcebox,masks,scale=3):
 im=Image.open(path).convert('RGB');d=ImageDraw.Draw(im)
 for r in masks:d.rectangle([(r[0]-sourcebox[0])*scale,(r[1]-sourcebox[1])*scale,(r[2]-sourcebox[0])*scale,(r[3]-sourcebox[1])*scale],fill='white')
 im.save(path)
for name,folder in [('male-anatomy','figures'),('ch14-label-male','labeling')]:
 clean_png(A/folder/(name+'.png'),[65,490,442,742],[[375,490,442,622]])
clean_png(A/'figures/female-anatomy.png',[158,405,539,742],[[158,405,539,421],[158,421,222,521]])
# Use the isolated source sperm picture rather than retaining part of a neighbouring panel.
from pptx import Presentation
sourceppt=ROOT/'authoring/source-inputs/Unit B Chapter 14 Notes.pptx'
sourceppt.parent.mkdir(parents=True,exist_ok=True)
if not sourceppt.exists():raise FileNotFoundError('Restore authoring/source-inputs/Unit B Chapter 14 Notes.pptx before rebuilding.')
prs=Presentation(sourceppt);picture=[s for s in prs.slides[10].shapes if s.shape_type==13][0]
(A/'figures/sperm-structure.png').write_bytes(picture.image.blob)
im=Image.open(A/'figures/sperm-structure.png').convert('RGB');d=ImageDraw.Draw(im)
for letter,rect in [('A',(154,32,236,59)),('B',(181,64,230,85)),('C',(63,135,127,180)),('D',(115,212,234,240))]:
 d.rectangle(rect,fill='white');d.text((rect[0]+2,rect[1]),letter,font=font,fill='black')
im.save(A/'labeling/ch14-label-sperm.png')
for r in records:
 if r['id'] in ['sperm-structure','ch14-label-sperm']:
  r['source']={'file':'authoring/source-inputs/Unit B Chapter 14 Notes.pptx','slide':11,'pictureIndex':0,'sha256':hashlib.sha256(sourceppt.read_bytes()).hexdigest()};r['changes']='Isolated source sperm figure. Four text labels covered/replaced in practice version; original structure and leader lines retained.'
  if r['id']=='sperm-structure':r['caption']='The acrosome, head, middle piece and flagellum have different roles.'
# Re-render the follicle crop with its final letter fully inside the canvas.
old=next(r for r in records if r['id']=='ch14-label-follicle');doc=fitz.open(P);p=doc[20]
for ch in old['changes']:
 r=fitz.Rect(ch['textCoverBox']);p.draw_rect(r,color=None,fill=(1,1,1));p.insert_text((r.x0+1,r.y1-1),ch['letter'],fontsize=11,fontname='hebo')
for r in [[160,490,342,524],[342,490,354,512],[230,710,357,739],[432,713,535,739]]:p.draw_rect(fitz.Rect(r),color=None,fill=(1,1,1))
p.get_pixmap(matrix=fitz.Matrix(3,3),clip=fitz.Rect(160,490,535,728)).save(A/'labeling/ch14-label-follicle.png');doc.close();old['crop']=[160,490,535,728]
relabel('ch14-label-female','female-anatomy',5,[158,422,539,742],[
 ('A','Ovary',[258,448,283,460]),('B','Oviduct',[324,433,357,445]),('C','Uterus',[164,535,194,547]),
 ('D','Fimbriae',[497,481,534,493]),('E','Cervix',[497,577,525,589]),('F','Vagina',[497,635,528,647])],
 'Female reproductive anatomy','Female reproductive anatomy with six source leader-line labels replaced by letters; other labels provide orientation.')
clean_png(A/'labeling/ch14-label-female.png',[158,422,539,742],[[158,422,222,521]])
(ROOT/'authoring/diagram-review.json').write_text(json.dumps(records,indent=2));(ROOT/'authoring/label-diagrams.json').write_text(json.dumps(labels,indent=2))
print('Final',len(labels),'labeling activities')
